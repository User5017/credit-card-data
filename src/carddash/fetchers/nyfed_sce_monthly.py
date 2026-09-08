"""New York Fed Survey of Consumer Expectations, the monthly core survey: one xlsx with the full history.

Publication facts, verified 2026-09-08:
- One workbook, always at the same URL, re-published every month with the whole history:
    https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-data.xlsx
  About 1.2 MB, downloads directly with this pipeline's own User-Agent, no redirect, no gate. No date in the name and
  no cache headers, so a new month shows up as a content diff under the canonical snapshot name.
- About 1,300 household heads a month, a rotating panel, from June 2013. Readings are dated to the survey month,
  so period_type is M and period_end the last day of that month. The monthly release lands about the second week
  of the following month.
- Every data sheet has the same shape: row 1 attribution, row 2 the sheet title, then one or two header rows, then
  one row per month with the date as an integer YYYYMM in column A. Months must be consecutive: a missing month
  fails the parse rather than thinning a line on the page.
- Loaded here, out of 45 sheets, the four that speak to card credit:
    'Delinquency expectations'       mean probability of missing a minimum debt payment in the next three months
    'Delinquency expectations Demo'  the same by age, education, income, numeracy and region (15 columns)
    'Credit availability'            shares saying credit is harder or easier than a year ago, and a year ahead
    'Household financial situation'  shares saying the household is worse or better off than a year ago, and a year ahead
  The two 'Year ago / Year ahead' sheets have a two-row header: the horizon on the upper row over the first cell of
  each block of five, the five answer categories on the lower row. The horizon is carried forward across its block.
- Values are percentages of respondents (weighted). Nothing is computed here; the shares are transcribed one per
  series, and sums like 'harder' = 'much harder' + 'somewhat harder' live in sql/views.sql.

Outputs: raw_dir/latest/frbny-sce-data.xlsx (byte for byte) and one facts row per month per series.
"""

from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ..schema import FACT_COLUMNS, last_day

SOURCE = "nyfed_sce_monthly"
LANDING_URL = "https://www.newyorkfed.org/microeconomics/sce"
FILE_URL = "https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-data.xlsx"
RAW_NAME = "frbny-sce-data.xlsx"
XLSX_MAGIC = b"PK\x03\x04"
PERIOD_TYPE = "M"
FIRST_MONTH = (2013, 6)

# demographic column header -> (entity, entity_type), for the 'Demo' sheets
DEMO_ENTITIES = {
    "Age Under 40": ("AGE:LT40", "age"),
    "Age 40-60": ("AGE:40-60", "age"),
    "Age Over 60": ("AGE:GT60", "age"),
    "Education High School or Less": ("EDU:HS_OR_LESS", "education"),
    "Education Some College": ("EDU:SOME_COLLEGE", "education"),
    "Education BA or Higher": ("EDU:BA_PLUS", "education"),
    "Income under 50k": ("INCOME:LT50K", "income"),
    "Income 50-100k": ("INCOME:50-100K", "income"),
    "Income Over 100k": ("INCOME:GT100K", "income"),
    "Numeracy Low": ("NUMERACY:LOW", "numeracy"),
    "Numeracy High": ("NUMERACY:HIGH", "numeracy"),
    "Region West": ("REGION:WEST", "region"),
    "Region Midwest": ("REGION:MIDWEST", "region"),
    "Region South": ("REGION:SOUTH", "region"),
    "Region Northeast": ("REGION:NORTHEAST", "region"),
}

# answer category -> metric suffix, for the two 'Year ago / Year ahead' sheets
CREDIT_ANSWERS = {
    "Much harder": "much_harder",
    "Somewhat harder": "somewhat_harder",
    "Equally easy/hard": "same",
    "Somewhat easier": "somewhat_easier",
    "Much easier": "much_easier",
}
FINANCE_ANSWERS = {
    "Much worse off": "much_worse",
    "Somewhat worse off": "somewhat_worse",
    "About the same": "same",
    "Somewhat better off": "somewhat_better",
    "Much better off": "much_better",
}
HORIZONS = {"Year ago": "year_ago", "Year ahead": "year_ahead"}

# sheet -> how to read it
SHEETS = {
    "Delinquency expectations": {"kind": "single", "metric": "sce_miss_payment_prob",
                                 "title": "Mean probability of missing minimum debt payment over the next three months"},
    "Delinquency expectations Demo": {"kind": "demo", "metric": "sce_miss_payment_prob"},
    "Credit availability": {"kind": "horizon", "prefix": "sce_credit", "answers": CREDIT_ANSWERS},
    "Household financial situation": {"kind": "horizon", "prefix": "sce_finances", "answers": FINANCE_ANSWERS},
}

_DATE_RE = re.compile(r"^(\d{4})(\d{2})$")


def looks_like_xlsx(content: bytes) -> bool:
    return content[:4] == XLSX_MAGIC


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell).strip()) if cell is not None else ""


def parse_month(value) -> tuple[int, int]:
    """201306 -> (2013, 6)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    m = _DATE_RE.match(str(value).strip())
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise ValueError(f"not a YYYYMM month: {value!r}")
    return int(m.group(1)), int(m.group(2))


def next_month(year: int, month: int) -> tuple[int, int]:
    return (year + 1, 1) if month == 12 else (year, month + 1)


def _num(value, where: str) -> float:
    bad = ValueError(f"{where}: expected a number, got {value!r}")
    if value is None or isinstance(value, bool) or isinstance(value, str):
        raise bad
    x = float(value)
    if not math.isfinite(x):
        raise bad
    return x


def read_sheet(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path.name}: no sheet {sheet!r}. Sheets: {wb.sheetnames}")
        return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def _split(rows: list[tuple], name: str) -> tuple[list[tuple], list[tuple]]:
    """(header rows, month rows): the month rows are the ones whose column A parses as YYYYMM."""
    first = next((i for i, r in enumerate(rows) if r and _DATE_RE.match(str(r[0]).strip().replace(".0", ""))), None)
    if first is None:
        raise ValueError(f"{name}: no month rows")
    data = [r for r in rows[first:] if r and r[0] is not None and str(r[0]).strip()]
    return rows[:first], data


def _columns(sheet: str, header_rows: list[tuple], name: str) -> dict[int, tuple[str, str, str]]:
    """column index -> (metric, entity, entity_type) for the sheet's kind."""
    spec = SHEETS[sheet]
    if spec["kind"] == "single":
        titles = [_norm(c) for r in header_rows for c in r if _norm(c)]
        if spec["title"] not in titles:
            raise ValueError(f"{name}: title {spec['title']!r} not found in {titles}")
        return {1: (spec["metric"], "SCE_ALL", "aggregate")}
    header = header_rows[-1]
    if spec["kind"] == "demo":
        out = {}
        for j, cell in enumerate(header):
            key = _norm(cell)
            if not key:
                continue
            if key not in DEMO_ENTITIES:
                raise ValueError(f"{name}: unknown demographic column {key!r}")
            entity, entity_type = DEMO_ENTITIES[key]
            out[j] = (spec["metric"], entity, entity_type)
        missing = set(DEMO_ENTITIES) - {_norm(c) for c in header}
        if missing:
            raise ValueError(f"{name}: demographic columns missing: {sorted(missing)}")
        return out
    # horizon: the row above carries 'Year ago' / 'Year ahead' over the first cell of each block
    upper = header_rows[-2]
    horizon = None
    out = {}
    for j, cell in enumerate(header):
        key = _norm(cell)
        if j < len(upper) and _norm(upper[j]) in HORIZONS:
            horizon = HORIZONS[_norm(upper[j])]
        if not key:
            continue
        if key not in spec["answers"]:
            raise ValueError(f"{name}: unknown answer column {key!r}")
        if horizon is None:
            raise ValueError(f"{name}: answer column {key!r} before any horizon label")
        out[j] = (f"{spec['prefix']}_{horizon}_{spec['answers'][key]}", "SCE_ALL", "aggregate")
    expected = len(HORIZONS) * len(spec["answers"])
    if len(out) != expected:
        raise ValueError(f"{name}: expected {expected} answer columns, found {len(out)}")
    return out


def parse_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """One sheet -> long frame [metric, entity, entity_type, period_end, value]. Months must be consecutive."""
    name = f"{path.name} {sheet!r}"
    header_rows, data = _split(read_sheet(path, sheet), name)
    columns = _columns(sheet, header_rows, name)
    records = []
    prev = None
    for row in data:
        year, month = parse_month(row[0])
        if prev is not None and (year, month) != next_month(*prev):
            raise ValueError(f"{name}: months jump from {prev[0]}-{prev[1]:02d} to {year}-{month:02d}")
        prev = (year, month)
        period_end = pd.Timestamp(last_day(year, month))
        for j, (metric, entity, entity_type) in columns.items():
            value = _num(row[j] if j < len(row) else None, f"{name} {year}-{month:02d} column {j}")
            records.append({"metric": metric, "entity": entity, "entity_type": entity_type,
                            "period_end": period_end, "value": value})
    if not records:
        raise ValueError(f"{name}: no rows")
    return pd.DataFrame.from_records(records)


def parse_release(path: Path, pulled_at: str) -> pd.DataFrame:
    df = pd.concat([parse_sheet(path, sheet) for sheet in SHEETS], ignore_index=True)
    dupes = df[df.duplicated(["metric", "entity", "period_end"], keep=False)]
    if not dupes.empty:
        raise ValueError(f"{path.name}: {len(dupes)} duplicate (metric, entity, month) rows")
    first = df["period_end"].min()
    if first != pd.Timestamp(last_day(*FIRST_MONTH)):
        raise ValueError(f"{path.name}: history starts {first.date()}, expected {last_day(*FIRST_MONTH)}")
    out = pd.DataFrame(
        {
            "metric": df["metric"],
            "entity": df["entity"],
            "entity_type": df["entity_type"],
            "tier": "all",
            "period_end": df["period_end"],
            "period_type": PERIOD_TYPE,
            "value": df["value"].astype("float64"),
            "source": SOURCE,
            "pulled_at": pulled_at,
        }
    )
    return out[FACT_COLUMNS]


def fetch(meta: pd.DataFrame, raw_dir: Path, session, pulled_at: str) -> pd.DataFrame:
    """Series come out of the sheets, not out of meta. The loader rejects any that series.csv does not list."""
    resp = session.get(FILE_URL)
    resp.raise_for_status()
    if not looks_like_xlsx(resp.content):
        raise ValueError(f"{FILE_URL}: not an xlsx (starts with {resp.content[:40]!r})")
    latest = raw_dir / "latest"
    latest.mkdir(parents=True, exist_ok=True)
    dest = latest / RAW_NAME
    dest.write_bytes(resp.content)
    return parse_release(dest, pulled_at)
