"""CFPB Terms of Credit Card Plans (TCCP) fetcher: a semiannual xlsx with one row per card product.

Publication facts, verified 2026-09-07:
- The survey page links one workbook per half-year on files.consumerfinance.gov. File names are irregular
  (cfpb_tccp-data_2025-12-31.xlsx, cfpb_tccp-data_2023-07-01_2023_12-31.xlsx, and a re-upload named
  cfpb_tccp-data_2024-12-31_uYJ9Krf.xlsx), so the fetcher reads the links off the page and dates each file by the
  last date in its name. Snapshots are saved under a canonical name so a re-upload shows up as a content change.
- The CFPB's Akamai edge answers 403 to a browser-like User-Agent sent from a script and 200 to this pipeline's own
  'carddash/<version> (+repo url)' agent (src/carddash/http.py). Do not fake a browser here.
- One sheet per workbook (its name varies). Rows 1-9 are a title block including
  'Data for reporting period ... - December 31, 2025', the header is row 10, column A is empty, then one row per
  product. The header row is found by content and the title-block end date must agree with the file name.
- Header names drift between vintages ('Periodic max' / 'Periodic Max', 'Tier 1 To Balance' / 'to Balance').
  'Issued by Top 25 Institution' exists from H2 2023, the 'no score' APR columns from H2 2023, 'Institution Type'
  only in the 2024 H2 re-upload. Optional columns come back as nulls.
- 'Report Date' is unreliable ('Data as of June 30' in H1 2025, with six rows saying December 31), so it is kept
  as text and never used for dating.
- APRs are fractions as published (0.2749 is 27.49%). 9.99 and 0 appear as placeholders. That, and every other
  aggregation choice, lives in sql/tccp_facts.sql. The parser only transcribes.
- H2 2022 and earlier use the pre-2023 survey layout (different columns, no tier medians) and are not loaded.
  The 1990-2022 archive on data.consumerfinance.gov is a third layout.

Outputs: raw_dir/latest/cfpb_tccp-data_<period_end>.xlsx (one per half-year), raw_dir/tccp_products.csv (the
sub-grain raw table, one row per product per half-year), and facts rows computed from it by sql/tccp_facts.sql.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

from ..paths import REPO_ROOT
from ..schema import FACT_COLUMNS
from ..schema import period_end as _half_year_end

SOURCE = "tccp"
SURVEY_URL = "https://www.consumerfinance.gov/data-research/credit-card-data/terms-credit-card-plans-survey/"
FIRST_PERIOD_END = dt.date(2023, 6, 30)  # first workbook in the 2023 survey layout
PRODUCTS_CSV = "tccp_products.csv"
SNAPSHOT_GLOB = "cfpb_tccp-data_*.xlsx"
FACTS_SQL = REPO_ROOT / "sql" / "tccp_facts.sql"

_LINK_RE = re.compile(r"https?://files\.consumerfinance\.gov/f/documents/cfpb_tccp-data_[A-Za-z0-9_\-]+\.xlsx")
_DATE_RE = re.compile(r"(\d{4})[-_](\d{2})[-_](\d{2})")
_TITLE_DATE_RE = re.compile(r"[A-Z][a-z]+ \d{1,2}, \d{4}")
_HEADER_MIN_CELLS = 20
_HEADER_SEARCH_ROWS = 30

# (output column, accepted header texts after normalization, kind, required)
# Optional columns are null in vintages that lack them. Everything else in the workbook stays in the raw snapshot.
COLUMN_SPEC = [
    ("institution", ("institution name",), "text", True),
    ("top25", ("issued by top 25 institution",), "bool", False),
    ("institution_type", ("institution type",), "text", False),
    ("product", ("product name", "card name"), "text", True),
    ("report_date", ("report date",), "text", False),
    ("availability", ("availability of credit card plan", "availability of credit", "offering location"), "text", False),
    ("secured", ("secured card",), "bool", True),
    ("targeted_tiers", ("targeted credit tiers",), "text", True),
    ("purchase_apr_offered", ("purchase apr offered?",), "bool", True),
    ("purchase_apr_varies_by_tier", ("purchase apr vary by credit tier",), "bool", True),
    ("apr_index", ("index",), "text", False),
    ("variable_rate_index", ("variable rate index",), "text", False),
    ("purchase_apr_no_score", ("purchase apr no score",), "num", False),
    ("purchase_apr_le619", ("purchase apr poor",), "num", True),
    ("purchase_apr_620_719", ("purchase apr good",), "num", True),
    ("purchase_apr_ge720", ("purchase apr great",), "num", True),
    ("purchase_apr_min", ("purchase apr min",), "num", True),
    ("purchase_apr_median", ("purchase apr median",), "num", True),
    ("purchase_apr_max", ("purchase apr max",), "num", True),
    ("annual_fee", ("annual fee",), "num", False),
    ("late_fee", ("late fee ($)",), "num", False),
]
# Boolean flags split out of the multi-select 'Targeted Credit Tiers' text (null when the text is empty).
TIER_FLAGS = [
    ("targets_no_score", "no credit score"),
    ("targets_le619", "619 or less"),
    ("targets_620_719", "620 to 719"),
    ("targets_ge720", "720 or greater"),
]
PRODUCT_COLUMNS = [
    "period_end", "source_file", "row_in_file",
    "institution", "top25", "institution_type", "product", "report_date", "availability", "secured",
    "targeted_tiers", "targets_no_score", "targets_le619", "targets_620_719", "targets_ge720",
    "purchase_apr_offered", "purchase_apr_varies_by_tier", "apr_index", "variable_rate_index",
    "purchase_apr_no_score", "purchase_apr_le619", "purchase_apr_620_719", "purchase_apr_ge720",
    "purchase_apr_min", "purchase_apr_median", "purchase_apr_max", "annual_fee", "late_fee",
]
_KINDS = {
    "period_end": "date", "source_file": "text", "row_in_file": "int",
    **{col: kind for col, _, kind, _ in COLUMN_SPEC},
    **{col: "bool" for col, _ in TIER_FLAGS},
}
_PANDAS_DTYPES = {"text": "string", "int": "Int64", "bool": "boolean", "num": "Float64"}
_DUCKDB_TYPES = {"date": "DATE", "text": "VARCHAR", "int": "INTEGER", "bool": "BOOLEAN", "num": "DOUBLE"}
PRODUCT_TYPES = {col: _DUCKDB_TYPES[_KINDS[col]] for col in PRODUCT_COLUMNS}  # schema for reading the CSV back


# ---------- discovery ----------


def period_end_from_name(name: str) -> dt.date:
    """'cfpb_tccp-data_2023-07-01_2023_12-31.xlsx' -> 2023-12-31. The last date wins, '_' and '-' both separate."""
    dates = _DATE_RE.findall(name)
    if not dates:
        raise ValueError(f"no date in TCCP file name {name!r}")
    y, m, d = (int(x) for x in dates[-1])
    pe = dt.date(y, m, d)
    if pe != _half_year_end(pe, "H"):
        raise ValueError(f"{name!r}: {pe} is not the end of a half-year")
    return pe


def snapshot_name(pe: dt.date) -> str:
    return f"cfpb_tccp-data_{pe.isoformat()}.xlsx"


def parse_index(html: str) -> dict[dt.date, str]:
    """Survey page HTML -> {period_end: workbook url} for every vintage linked, whatever the file is called."""
    found: dict[dt.date, str] = {}
    for url in dict.fromkeys(_LINK_RE.findall(html)):
        pe = period_end_from_name(url.rsplit("/", 1)[-1])
        if pe in found:
            raise ValueError(f"two TCCP workbooks for {pe}: {found[pe]} and {url}")
        found[pe] = url
    if not found:
        raise ValueError("no TCCP workbook links on the survey page (layout changed, or the request was blocked)")
    return found


def download_workbook(session, url: str, dest: Path) -> None:
    resp = session.get(url)
    resp.raise_for_status()
    if not resp.content.startswith(b"PK"):
        raise ValueError(f"{url}: not an xlsx (starts with {resp.content[:40]!r})")
    dest.write_bytes(resp.content)


# ---------- parsing ----------


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell).strip().lower()) if cell is not None else ""


def _text(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def _bool(v, col: str, row: int):
    s = _norm(v)
    if not s:
        return None
    if s in ("yes", "true"):
        return True
    if s in ("no", "false"):
        return False
    raise ValueError(f"row {row}, {col}: expected Yes/No or True/False, got {v!r}")


def _num(v, col: str, row: int):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError(f"row {row}, {col}: expected a number, got {v!r}")
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except ValueError:
        raise ValueError(f"row {row}, {col}: expected a number, got {v!r}") from None


def _read_rows(path: Path) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) != 1:
            raise ValueError(f"{path.name}: expected one sheet, found {wb.sheetnames}")
        return [tuple(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    finally:
        wb.close()


def _header_index(rows: list[tuple]) -> int:
    for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        if sum(1 for c in row if _norm(c)) >= _HEADER_MIN_CELLS:
            return i
    raise ValueError(f"no header row (one with {_HEADER_MIN_CELLS}+ named cells) in the first {_HEADER_SEARCH_ROWS} rows")


def _title_period_end(rows: list[tuple], header_index: int) -> dt.date:
    for row in rows[:header_index]:
        for c in row:
            if isinstance(c, str) and "reporting period" in c.lower():
                dates = _TITLE_DATE_RE.findall(c)
                if not dates:
                    raise ValueError(f"title block line without a date: {c!r}")
                return dt.datetime.strptime(dates[-1], "%B %d, %Y").date()
    raise ValueError("title block has no 'reporting period' line above the header")


def _column_positions(header: tuple, name: str) -> dict[str, int]:
    """Output column -> cell index, from header text. Required columns must exist, used ones must be unique."""
    positions: dict[str, list[int]] = {}
    for j, cell in enumerate(header):
        key = _norm(cell)
        if key:
            positions.setdefault(key, []).append(j)
    out: dict[str, int] = {}
    missing = []
    for col, aliases, _, required in COLUMN_SPEC:
        idx = [j for a in aliases for j in positions.get(a, [])]
        if not idx:
            if required:
                missing.append(aliases[0])
            continue
        if len(idx) > 1:
            raise ValueError(f"{name}: header {aliases[0]!r} appears {len(idx)} times")
        out[col] = idx[0]
    if missing:
        raise ValueError(f"{name}: required headers missing: {missing}. Found: {sorted(positions)}")
    return out


def _frame(records: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame.from_records(records, columns=PRODUCT_COLUMNS)
    for col in PRODUCT_COLUMNS:
        kind = _KINDS[col]
        df[col] = pd.to_datetime(df[col]) if kind == "date" else df[col].astype(_PANDAS_DTYPES[kind])
    return df


def parse_workbook(path: Path, period_end: dt.date) -> pd.DataFrame:
    """One TCCP workbook -> product rows (PRODUCT_COLUMNS). Header text drives everything, never positions."""
    rows = _read_rows(path)
    h = _header_index(rows)
    title_end = _title_period_end(rows, h)
    if title_end != period_end:
        raise ValueError(f"{path.name}: title block says the period ends {title_end}, the file name says {period_end}")
    pos = _column_positions(rows[h], path.name)
    records = []
    for i, row in enumerate(rows[h + 1 :], start=h + 2):  # i is the 1-based row number in the sheet
        if not any(_norm(c) for c in row):
            continue
        rec: dict = {"period_end": period_end, "source_file": path.name, "row_in_file": i}
        for col, _, kind, _ in COLUMN_SPEC:
            j = pos.get(col)
            v = row[j] if j is not None and j < len(row) else None
            rec[col] = _text(v) if kind == "text" else _bool(v, col, i) if kind == "bool" else _num(v, col, i)
        if rec["institution"] is None:
            raise ValueError(f"{path.name} row {i}: has content but no institution name")
        tiers = (rec["targeted_tiers"] or "").lower()
        for flag, needle in TIER_FLAGS:
            rec[flag] = (needle in tiers) if tiers else None
        if tiers and not any(rec[flag] for flag, _ in TIER_FLAGS):
            raise ValueError(f"{path.name} row {i}: unrecognized targeted credit tiers {rec['targeted_tiers']!r}")
        records.append(rec)
    if not records:
        raise ValueError(f"{path.name}: no product rows under the header")
    return _frame(records)


def parse_all(latest_dir: Path) -> pd.DataFrame:
    """Every snapshot in raw_dir/latest -> one product table, dated by file name."""
    paths = sorted(latest_dir.glob(SNAPSHOT_GLOB))
    if not paths:
        raise ValueError(f"no TCCP snapshots in {latest_dir}")
    return pd.concat([parse_workbook(p, period_end_from_name(p.name)) for p in paths], ignore_index=True)


def write_products(products: pd.DataFrame, path: Path) -> None:
    out = products.sort_values(["period_end", "row_in_file"]).reset_index(drop=True)
    out = out.assign(period_end=out["period_end"].dt.strftime("%Y-%m-%d"))
    out.to_csv(path, index=False, lineterminator="\n", float_format="%.10g")


# ---------- facts ----------


def _read_sql(path: Path) -> str:
    return "\n".join(line for line in path.read_text(encoding="utf-8").splitlines() if not line.lstrip().startswith("--"))


def products_to_facts(products: pd.DataFrame, pulled_at: str) -> pd.DataFrame:
    """Run sql/tccp_facts.sql over the product table. The SQL owns every aggregation choice."""
    con = duckdb.connect()
    try:
        con.register("tccp_products", products)
        out = con.execute(_read_sql(FACTS_SQL)).df()
    finally:
        con.close()
    if out.empty:
        raise ValueError("sql/tccp_facts.sql produced no rows")
    df = pd.DataFrame(
        {
            "metric": out["metric"],
            "entity": out["entity"],
            "entity_type": "aggregate",
            "tier": out["tier"],
            "period_end": pd.to_datetime(out["period_end"]),
            "period_type": "H",
            "value": out["value"].astype("float64"),
            "source": SOURCE,
            "pulled_at": pulled_at,
        }
    )
    return df[FACT_COLUMNS]


def fetch(meta: pd.DataFrame, raw_dir: Path, session, pulled_at: str) -> pd.DataFrame:
    """Series come out of the SQL, not out of meta. The loader rejects any that series.csv does not list."""
    latest = raw_dir / "latest"
    latest.mkdir(parents=True, exist_ok=True)
    resp = session.get(SURVEY_URL)
    resp.raise_for_status()
    links = {pe: url for pe, url in parse_index(resp.text).items() if pe >= FIRST_PERIOD_END}
    if not links:
        raise ValueError(f"the survey page lists no workbook from {FIRST_PERIOD_END} on")
    keep = set()
    for pe in sorted(links):
        dest = latest / snapshot_name(pe)
        download_workbook(session, links[pe], dest)
        keep.add(dest.name)
    for stale in latest.glob(SNAPSHOT_GLOB):  # latest/ mirrors the page, as replace-by-source does for facts
        if stale.name not in keep:
            stale.unlink()
    products = parse_all(latest)
    write_products(products, raw_dir / PRODUCTS_CSV)
    return products_to_facts(products, pulled_at)
