"""New York Fed Survey of Consumer Expectations, Credit Access Survey: one xlsx with the full history.

Publication facts, verified 2026-09-08:
- One workbook, always at the same URL, re-published with every wave:
    https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-credit-access-data.xlsx
  It downloads directly with this pipeline's own User-Agent (107 KB). The landing page shows a confirm-download
  modal, but that is client side only: there is no server-side gate, no cookie and no terms click. The response
  carries neither Last-Modified nor ETag, so a new wave can only be detected by content, which is why the snapshot
  is saved under a canonical name and the loader's diff does the detecting.
- Waves run every four months, fielded in February, June and October, from October 2013. The credit access module
  is asked of the SCE panel; each reading covers the twelve months up to and including the survey month, so a
  reading is dated to the last day of the survey month and its period_type is T (every four months). A wave lands
  about four to six weeks after fielding.
- Four sheets: Disclaimer, License, overall, demographics. Both data sheets carry an attribution line in row 1, the
  header in row 2, then one row per (date, group, category). date is an integer YYYYMM. Values are percentages of
  respondents, already weighted; Observations is the unweighted respondent count.
- 'overall' has 36 columns and one group (all/Overall). It is the only sheet with the card-specific rejection
  questions: CCRejected, CCLimitRejected and the expectation ChanceCCApplicationWillBeRejected.
- 'demographics' has 26 columns and six groups: credit_score (less_680, between_680_760, over_760, from February
  2014) and age (less_eq_40, between_40_59, over_60). It carries application and limit-request rates but not the
  card-specific rejection rates, so rejection by score is only available for credit of any kind.
- Sample sizes are small in the tails: the sub-680 credit score bucket had 151 respondents in the June 2026 wave
  against 508 for the over-760 bucket. Observations is loaded as a series for every entity so the page can show
  what a reading rests on, and the scope notes say so.
- A question that was not asked in a wave is written 'N/A'. Exactly one cell is like that today: the expectation
  question ChanceCCApplicationWillBeRejected was not asked in the first wave, October 2013. Such a cell produces no
  fact row, but only where the gap is at the start of the series: a hole in the middle, or a question that stops
  being asked, fails the parse, because either would silently shorten a line on the page.
- The credit score buckets are the respondent's self-reported score band, not a bureau score, and they do not line
  up with the CFPB tiers in crosswalks/tiers.csv. They are kept as their own entities (SCORE:...) and never mapped.

Outputs: raw_dir/latest/frbny-sce-credit-access-data.xlsx (the workbook byte for byte, under its own name) and one
facts row per wave per series. Nothing is computed here; the survey's own percentages are transcribed.
"""

from __future__ import annotations

import datetime as dt
import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ..schema import FACT_COLUMNS, last_day

SOURCE = "nyfed_sce"
LANDING_URL = "https://www.newyorkfed.org/microeconomics/sce/credit-access"
FILE_URL = "https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-credit-access-data.xlsx"
RAW_NAME = "frbny-sce-credit-access-data.xlsx"
XLSX_MAGIC = b"PK\x03\x04"
PERIOD_TYPE = "T"  # every four months
SURVEY_MONTHS = (2, 6, 10)
HEADER_ROW = 1  # zero-based: row 1 is the attribution line, row 2 the header
FIRST_WAVE = (2013, 10)

# (sheet group, sheet category) -> (entity, entity_type)
ENTITIES = {
    ("all", "Overall"): ("SCE_ALL", "aggregate"),
    ("credit_score", "less_680"): ("SCORE:LT680", "score"),
    ("credit_score", "between_680_760"): ("SCORE:680-760", "score"),
    ("credit_score", "over_760"): ("SCORE:GE760", "score"),
    ("age", "less_eq_40"): ("AGE:LE40", "age"),
    ("age", "between_40_59"): ("AGE:40-59", "age"),
    ("age", "over_60"): ("AGE:GE60", "age"),
}
# column header in either sheet -> metric. Headers are matched after whitespace and case normalization.
SHARED_COLUMNS = {
    "AppliedforCreditCard": "sce_card_application_rate",
    "RequestedIncreaseinCreditCardLimit": "sce_card_limit_request_rate",
    "Appliedforoneormoretypesofcredit": "sce_any_application_rate",
    "Rejectedfromoneormoresourcesofcredit_givenapplied": "sce_any_rejection_rate",
    "Discouraged": "sce_discouraged_rate",
    "Lenderclosedatleastoneacct": "sce_lender_closed_rate",
    "Observations": "sce_observations",
}
# only the overall sheet asks these
OVERALL_ONLY_COLUMNS = {
    "CCRejected": "sce_card_rejection_rate",
    "CCLimitRejected": "sce_card_limit_rejection_rate",
    "ChanceCCApplicationWillBeRejected": "sce_card_expected_rejection_rate",
}
SHEETS = {"overall": SHARED_COLUMNS | OVERALL_ONLY_COLUMNS, "demographics": SHARED_COLUMNS}
COUNT_METRICS = {"sce_observations"}
NOT_ASKED = {"n/a", "na", "."}  # how the workbook marks a question that was not put in that wave

_DATE_RE = re.compile(r"^(\d{4})(\d{2})$")


# ---------- waves ----------


def parse_wave(value) -> tuple[int, int]:
    """201310 -> (2013, 10). Anything that is not a survey month in YYYYMM fails."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    m = _DATE_RE.match(str(value).strip())
    if not m:
        raise ValueError(f"not a YYYYMM wave: {value!r}")
    year, month = int(m.group(1)), int(m.group(2))
    if month not in SURVEY_MONTHS:
        raise ValueError(f"wave {value} is month {month}, the survey is fielded in {SURVEY_MONTHS}")
    return year, month


def wave_end(year: int, month: int) -> dt.date:
    """A reading is dated to the last day of the month it was fielded in."""
    return last_day(year, month)


def next_wave(year: int, month: int) -> tuple[int, int]:
    return (year + 1, 2) if month == 10 else (year, month + 4)


# ---------- parsing ----------


def _norm(cell) -> str:
    return re.sub(r"\s+", "", str(cell).strip()).lower() if cell is not None else ""


def looks_like_xlsx(content: bytes) -> bool:
    return content[:4] == XLSX_MAGIC


def _num(value, header: str, where: str) -> float | None:
    """A numeric cell -> float. A documented not-asked marker -> None. Anything else fails."""
    bad = ValueError(f"{where}, {header!r}: expected a number, got {value!r}")
    if isinstance(value, str) and value.strip().lower() in NOT_ASKED:
        return None
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        raise bad
    try:
        x = float(value) if isinstance(value, (int, float)) else float(str(value).strip().replace(",", ""))
    except ValueError:
        raise bad from None
    if not math.isfinite(x):
        raise bad
    return x


def read_sheet(path: Path, sheet: str) -> tuple[list, list[tuple]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path.name}: no sheet {sheet!r}. Sheets: {wb.sheetnames}")
        rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()
    if len(rows) <= HEADER_ROW + 1:
        raise ValueError(f"{path.name} {sheet!r}: no data rows under the header")
    return list(rows[HEADER_ROW]), rows[HEADER_ROW + 1 :]


def _column_positions(header: list, wanted: dict[str, str], name: str) -> dict[str, int]:
    positions: dict[str, list[int]] = {}
    for j, cell in enumerate(header):
        key = _norm(cell)
        if key:
            positions.setdefault(key, []).append(j)
    for required in ("date", "group", "category"):
        if len(positions.get(required, [])) != 1:
            raise ValueError(f"{name}: column {required!r} appears {len(positions.get(required, []))} times")
    out = {"date": positions["date"][0], "group": positions["group"][0], "category": positions["category"][0]}
    missing = []
    for column in wanted:
        idx = positions.get(_norm(column), [])
        if len(idx) != 1:
            missing.append(f"{column} ({len(idx)} matches)")
            continue
        out[column] = idx[0]
    if missing:
        raise ValueError(f"{name}: columns missing or ambiguous: {missing}. Found: {sorted(positions)}")
    return out


def parse_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """One data sheet -> long frame [entity, entity_type, metric, period_end, value].

    Every row must name a known group and category, every wave must be a survey month, and each entity's waves must
    step by four months with no gap: a dropped wave or a renamed bucket fails here rather than thinning a series.
    """
    header, rows = read_sheet(path, sheet)
    wanted = SHEETS[sheet]
    name = f"{path.name} {sheet!r}"
    pos = _column_positions(header, wanted, name)
    seen: dict[str, list[tuple[int, int]]] = {}
    asked: dict[tuple[str, str], list] = {}
    records = []
    for i, row in enumerate(rows, start=HEADER_ROW + 2):  # 1-based sheet row number
        raw_date = row[pos["date"]] if pos["date"] < len(row) else None
        if raw_date is None or not str(raw_date).strip():
            if any(str(c).strip() for c in row if c is not None):
                raise ValueError(f"{name} row {i}: values with no date")
            continue
        year, month = parse_wave(raw_date)
        group = str(row[pos["group"]]).strip() if pos["group"] < len(row) else ""
        category = str(row[pos["category"]]).strip() if pos["category"] < len(row) else ""
        key = ENTITIES.get((group, category))
        if key is None:
            raise ValueError(f"{name} row {i}: unknown group/category {(group, category)}. Known: {sorted(ENTITIES)}")
        entity, entity_type = key
        prior = seen.setdefault(entity, [])
        if prior and (year, month) != next_wave(*prior[-1]):
            raise ValueError(
                f"{name} row {i}: {entity} jumps from {prior[-1][0]}-{prior[-1][1]:02d} to {year}-{month:02d}, "
                "waves must step by four months"
            )
        prior.append((year, month))
        period_end = pd.Timestamp(wave_end(year, month))
        for column, metric in wanted.items():
            value = _num(row[pos[column]] if pos[column] < len(row) else None, column, f"{name} row {i}")
            asked.setdefault((metric, entity), []).append((year, month, value))
            if value is not None:
                records.append(
                    {
                        "metric": metric,
                        "entity": entity,
                        "entity_type": entity_type,
                        "period_end": period_end,
                        "value": value,
                    }
                )
    if not records:
        raise ValueError(f"{name}: no wave rows")
    _check_not_asked(asked, name)
    return pd.DataFrame.from_records(records)


def _check_not_asked(asked: dict[tuple[str, str], list], name: str) -> None:
    """A not-asked marker is only allowed at the start of a series: the question was added later. A gap in the
    middle, or a question that stops being asked, would quietly shorten a line and must fail instead."""
    for (metric, entity), waves in asked.items():
        present = [i for i, (_, _, v) in enumerate(waves) if v is not None]
        if not present:
            raise ValueError(f"{name}: {metric} for {entity} is not asked in any wave")
        first = present[0]
        holes = [waves[i] for i in range(first, len(waves)) if waves[i][2] is None]
        if holes:
            where = ", ".join(f"{y}-{m:02d}" for y, m, _ in holes[:4])
            raise ValueError(
                f"{name}: {metric} for {entity} has no value in {where} after being asked from "
                f"{waves[first][0]}-{waves[first][1]:02d}; a mid-series gap is not a late-added question"
            )


def parse_release(path: Path, pulled_at: str) -> pd.DataFrame:
    """The workbook -> facts rows. The overall sheet also carries the card-specific rejection questions."""
    frames = [parse_sheet(path, sheet) for sheet in SHEETS]
    df = pd.concat(frames, ignore_index=True)
    dupes = df[df.duplicated(["metric", "entity", "period_end"], keep=False)]
    if not dupes.empty:
        raise ValueError(f"{path.name}: {len(dupes)} duplicate (metric, entity, wave) rows")
    first = df["period_end"].min()
    if first != pd.Timestamp(wave_end(*FIRST_WAVE)):
        raise ValueError(f"{path.name}: history starts {first.date()}, expected {wave_end(*FIRST_WAVE)}")
    out = pd.DataFrame(
        {
            "metric": df["metric"],
            "entity": df["entity"],
            "entity_type": df["entity_type"],
            "tier": "all",
            "period_end": df["period_end"],
            "period_type": PERIOD_TYPE,
            "value": df["value"].astype("float64"),
            "source": SOURCE,
            "pulled_at": pulled_at,
        }
    )
    return out[FACT_COLUMNS]


# ---------- fetch ----------


def fetch(meta: pd.DataFrame, raw_dir: Path, session, pulled_at: str) -> pd.DataFrame:
    """Series come out of the sheets, not out of meta. The loader rejects any that series.csv does not list."""
    resp = session.get(FILE_URL)
    resp.raise_for_status()
    if not looks_like_xlsx(resp.content):
        raise ValueError(f"{FILE_URL}: not an xlsx (starts with {resp.content[:40]!r})")
    latest = raw_dir / "latest"
    latest.mkdir(parents=True, exist_ok=True)
    dest = latest / RAW_NAME
    dest.write_bytes(resp.content)
    return parse_release(dest, pulled_at)
