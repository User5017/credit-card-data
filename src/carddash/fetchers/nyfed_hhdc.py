"""New York Fed Quarterly Report on Household Debt and Credit (HHDC) fetcher: one xlsx per release, full history.

Publication facts, verified 2026-09-07:
- One workbook per quarterly release:
    https://www.newyorkfed.org/medialibrary/interactives/householdcredit/data/xls/HHD_C_Report_2026Q2.xlsx
  The file name embeds the DATA quarter (2026Q2 was released 2026-08-11). It downloads directly with this pipeline's
  own User-Agent; there is no terms-of-use click (the landing page only has a cookie banner). A missing quarter
  redirects to /errors/404, which is served as HTTP 200 text/html, so discovery checks the body (an xlsx starts with
  the zip signature PK\\x03\\x04), never the status code. Files back to at least 2019Q1 stay online.
- Every workbook carries the full history (2003Q1 on for the loan-type sheets, 2000Q1 on for the age sheets), so
  only the newest release is downloaded. Prior quarters do get revised: the 2026Q2 file carries a footnote that
  2026Q1 card balances were revised. The loader's revision log covers that.
- Layout: a 'TABLE OF CONTENTS' sheet, then a 'ChartN' / 'Page N Data' pair per chart of the PDF report. Page
  numbers shift when charts are added, so data sheets are found by the title in cell A1, never by sheet name.
  Each data sheet has the title in A1, the unit in A2 ('Trillions of $', 'Millions', 'Percent'), a few navigation,
  source and note rows, then a header row (column A empty or 'quarter') and one row per quarter labelled 'YY:Qn'
  ('03:Q1' is 2003 Q1).
- Quirks handled on purpose: the account-count sheet labels 2003Q1-2011Q1 with date cells (the first day of the
  quarter's last month, 2003-03-01) before switching to 'YY:Qn'; the credit-limit sheet interleaves an unlabelled
  row of HE Revolving values after every quarter row; the balance sheet ends with a revision footnote; the two
  delinquency-flow sheets carry the header on row 5, not row 4. Quarters must be consecutive within a sheet. Any
  row that is not a quarter but has a value in a column we read fails the parse, and so does an empty or
  non-numeric cell in a quarter row: every read column is full in the 2025Q2, 2026Q1 and 2026Q2 releases.
- 'Credit Card' in this report means bankcards on credit reports; retail cards sit in 'Other' (report PDF
  HHDC_2026Q2.pdf p. 2: 'Other balances, which include retail cards and consumer finance loans', and the data
  dictionary on p. 45). The delinquency flows are four-quarter moving sums: balances newly 30+ (90+) days late in
  the quarter over balances that were current or less than 30 (90) days late the quarter before, an annualized
  share (press release footnote).
- Between the 2026Q1 and 2026Q2 releases the 2026Q1 card balance moved from 1,252 to 1,242 billion and limits by
  10 billion; the flows did not move, and nothing changed between the 2025Q2 and 2026Q2 files. Revisions are rare
  and large, so the newest-quarter balance golden entry is fixture-only and the flow entries are checked live.
- The limit sheet repeats the card balance to fewer decimals (up to $5 billion apart from 2009Q4 to 2012Q1, exact
  from 2012Q2 on); it is loaded once, from the balance sheet. The age sheet's 'all' column is computed over borrowers with a known birth year and
  differs from the loan-type sheet (6.9995 vs 6.97 for 2026Q2), so it is not loaded; the headline flow comes from
  the loan-type sheet. Card balances by age or state are not in the workbook (only total debt is split that way).
- Press releases: https://www.newyorkfed.org/newsevents/news/research/<year>/<yyyymmdd>. Report PDF:
  .../householdcredit/data/pdf/HHDC_<year>Q<n>.pdf. Suggested citation: Federal Reserve Bank of New York,
  Quarterly Report on Household Debt and Credit, Consumer Credit Panel/Equifax.

Outputs: raw_dir/latest/HHD_C_Report.xlsx (the release file byte for byte, under a canonical name so a new quarter
shows as a content diff) and one facts row per quarter per series in SHEET_SPEC. Trillions are turned into billions
here, with the factor sitting next to the unit text it is checked against; nothing else is computed.
"""

from __future__ import annotations

import datetime as dt
import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ..schema import FACT_COLUMNS, last_day

SOURCE = "nyfed_hhdc"
ENTITY = "CCP_ALL"  # the whole Consumer Credit Panel
LANDING_URL = "https://www.newyorkfed.org/microeconomics/hhdc"
FILE_URL = "https://www.newyorkfed.org/medialibrary/interactives/householdcredit/data/xls/HHD_C_Report_{year}Q{q}.xlsx"
RAW_NAME = "HHD_C_Report.xlsx"
MAX_WALK_BACK = 8  # quarters; the report lands five to seven weeks after quarter end
XLSX_MAGIC = b"PK\x03\x04"
TRILLIONS_TO_BILLIONS = 1000.0
AGE_GROUPS = ("18-29", "30-39", "40-49", "50-59", "60-69", "70+")

_DATA_SHEET_RE = re.compile(r"^Page \d+ Data$")
_LABEL_RE = re.compile(r"^(\d{2}):Q([1-4])$")
_QUARTER_END_MONTH = {1: 3, 2: 6, 3: 9, 4: 12}
_HEADER_SEARCH_ROWS = 10

# One entry per data sheet: the title in A1, the unit text in A2 (checked), the factor applied to every value in the
# sheet, and (header text, metric, entity, entity_type) per column read. Titles and headers match after whitespace
# normalization and case-folding. Columns not listed are ignored; listed ones must appear exactly once.
SHEET_SPEC = [
    {
        "title": "Total Debt Balance and Its Composition",
        "unit": "Trillions of $",
        "factor": TRILLIONS_TO_BILLIONS,
        "columns": [("Credit Card", "hhdc_card_balances", ENTITY, "aggregate")],
    },
    {
        "title": "Number of Accounts by Loan Type",
        "unit": "Millions",
        "factor": 1.0,
        "columns": [("Credit Card", "hhdc_card_accounts", ENTITY, "aggregate")],
    },
    {
        # 'Credit Card Balance' here repeats the balance sheet's column to fewer decimals (checked in the tests, not
        # loaded twice); 'Credit Card Available Credit' is limit minus balance and is not loaded either.
        "title": "Credit Limit and Balance for Credit Cards and HE Revolving",
        "unit": "Trillions of $",
        "factor": TRILLIONS_TO_BILLIONS,
        "columns": [("Credit Card Limit", "hhdc_card_limit", ENTITY, "aggregate")],
    },
    {
        "title": "Percent of Balance 90+ Days Delinquent by Loan Type",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [("CC", "hhdc_card_dq90_rate_balances", ENTITY, "aggregate")],
    },
    {
        "title": "New Delinquent* Balances by Loan Type",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [("CC", "hhdc_card_transition_dq30", ENTITY, "aggregate")],
    },
    {
        "title": "New Seriously Delinquent* Balances by Loan Type",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [("CC", "hhdc_card_transition_dq90", ENTITY, "aggregate")],
    },
    {
        "title": "Transition into Serious Delinquency (90+) for Credit Cards by Age",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [(g, "hhdc_card_transition_dq90", f"AGE:{g}", "age") for g in AGE_GROUPS],
    },
    # The four sheets below cover all consumer debt on the credit report, not cards alone (the workbook does not
    # split them by loan type). Their metric names say debt, accounts or consumers rather than card, and their scope
    # notes in series.csv say so.
    {
        "title": "Total Number of New and Closed Accounts and Consumer Credit Inquiries",
        "unit": "Millions",
        "factor": 1.0,
        "columns": [
            ("inquiry within 6 mo", "hhdc_inquiries_6mo", ENTITY, "aggregate"),
            ("closed within 12 mo", "hhdc_accounts_closed_12mo", ENTITY, "aggregate"),
            ("open within 12 mo", "hhdc_accounts_opened_12mo", ENTITY, "aggregate"),
        ],
    },
    {
        # the 'Total' column is a formula (=SUM) and is not read
        "title": "Total Balance by Delinquency Status",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [
            ("Current", "hhdc_debt_share_current", ENTITY, "aggregate"),
            ("30 days late", "hhdc_debt_share_dq30", ENTITY, "aggregate"),
            ("60 days late", "hhdc_debt_share_dq60", ENTITY, "aggregate"),
            ("90 days late", "hhdc_debt_share_dq90", ENTITY, "aggregate"),
            ("120+ days late", "hhdc_debt_share_dq120", ENTITY, "aggregate"),
            ("Severely Derogatory", "hhdc_debt_share_derogatory", ENTITY, "aggregate"),
        ],
    },
    {
        "title": "Number of Consumers with New Foreclosures and Bankruptcies",
        "unit": "Thousands",
        "factor": 1.0,
        "columns": [
            ("foreclosure", "hhdc_new_foreclosures", ENTITY, "aggregate"),
            ("bankruptcy", "hhdc_new_bankruptcies", ENTITY, "aggregate"),
        ],
    },
    {
        # A2 says Percent; the second column is dollars, which series.csv carries
        "title": "Third Party Collections",
        "unit": "Percent",
        "factor": 1.0,
        "columns": [
            ("proportion of consumers with collection", "hhdc_collections_share", ENTITY, "aggregate"),
            ("average collection amount per person with item", "hhdc_collections_avg_amount", ENTITY, "aggregate"),
        ],
    },
]


# ---------- quarters ----------


def quarter_of(date: dt.date) -> tuple[int, int]:
    return date.year, (date.month - 1) // 3 + 1


def quarter_end(year: int, q: int) -> dt.date:
    return last_day(year, _QUARTER_END_MONTH[q])


def previous_quarter(year: int, q: int) -> tuple[int, int]:
    return (year - 1, 4) if q == 1 else (year, q - 1)


def next_quarter(year: int, q: int) -> tuple[int, int]:
    return (year + 1, 1) if q == 4 else (year, q + 1)


def parse_quarter(label) -> tuple[int, int]:
    """'03:Q1' -> (2003, 1). A date cell (2003-03-01, how the account sheet labels 2003-2011) -> its quarter.

    Two-digit years are read as 2000 + YY; the consecutive-quarter check in parse_sheet catches a 1999 row.
    """
    if isinstance(label, dt.datetime):  # openpyxl hands date cells over as datetime
        return quarter_of(label.date())
    if isinstance(label, dt.date):
        return quarter_of(label)
    if isinstance(label, str):
        m = _LABEL_RE.match(label.strip())
        if m:
            return 2000 + int(m.group(1)), int(m.group(2))
    raise ValueError(f"not a quarter label: {label!r}")


def file_url(year: int, q: int) -> str:
    return FILE_URL.format(year=year, q=q)


# ---------- discovery ----------


def looks_like_xlsx(content: bytes) -> bool:
    """A missing file comes back as an HTML 404 page with status 200, so the body decides."""
    return content[:4] == XLSX_MAGIC


def discover_latest(session, today: dt.date, max_back: int = MAX_WALK_BACK) -> tuple[int, int, bytes]:
    """Newest quarter whose workbook exists: try the calendar quarter of the run date, walk back.

    Returns (year, quarter, workbook_bytes) so the file is downloaded once.
    """
    year, q = quarter_of(today)
    tried = []
    for _ in range(max_back + 1):
        resp = session.get(file_url(year, q))
        if resp.status_code == 404:  # not how the site answers today, but it would only mean 'not published yet'
            found = False
        else:
            resp.raise_for_status()  # anything else that is not 200 is a real outage (the session already retried)
            found = looks_like_xlsx(resp.content)
        if found:
            return year, q, resp.content
        tried.append(f"{year}Q{q}")
        year, q = previous_quarter(year, q)
    raise ValueError(f"no HHD_C_Report workbook found for any of {tried} (URL pattern changed?)")


# ---------- parsing ----------


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell).strip()).lower() if cell is not None else ""


def _cell(row: tuple, j: int):
    return row[j] if j < len(row) else None


def _num(v, header: str, name: str, row: int) -> float:
    """A numeric cell -> float. An empty, boolean, text or non-finite cell fails: a quarter row must carry a value.

    A silently dropped quarter would leave a 182-day gap, below the loader's continuity warning, so it must not
    be tolerated here.
    """
    bad = ValueError(f"{name} row {row}, {header!r}: expected a number, got {v!r}")
    if v is None or isinstance(v, bool) or (isinstance(v, str) and not v.strip()):
        raise bad
    try:
        x = float(v) if isinstance(v, (int, float)) else float(str(v).strip().replace(",", ""))
    except ValueError:
        raise bad from None
    if not math.isfinite(x):
        raise bad
    return x


def read_sheets(path: Path) -> dict[str, list[tuple]]:
    """Normalized A1 title -> rows, for every 'Page N Data' sheet. Titles must be unique."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out: dict[str, list[tuple]] = {}
        for sheet in wb.sheetnames:
            if not _DATA_SHEET_RE.match(sheet):
                continue
            rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
            title = _norm(_cell(rows[0], 0)) if rows else ""
            if not title:
                raise ValueError(f"{path.name}: sheet {sheet!r} has no title in A1")
            if title in out:
                raise ValueError(f"{path.name}: two data sheets titled {rows[0][0]!r}")
            out[title] = rows
        if not out:
            raise ValueError(f"{path.name}: no 'Page N Data' sheets; found {wb.sheetnames}")
        return out
    finally:
        wb.close()


def _header_row(rows: list[tuple], headers: list[str], name: str) -> int:
    wanted = {_norm(h) for h in headers}
    for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        if wanted <= {_norm(c) for c in row}:
            return i
    raise ValueError(f"{name}: headers missing: no row in the first {_HEADER_SEARCH_ROWS} has all of {headers}")


def _column_positions(header_row: tuple, headers: list[str], name: str) -> dict[str, int]:
    positions: dict[str, list[int]] = {}
    for j, cell in enumerate(header_row):
        key = _norm(cell)
        if key:
            positions.setdefault(key, []).append(j)
    out = {}
    for h in headers:
        idx = positions.get(_norm(h), [])
        if len(idx) != 1:
            raise ValueError(
                f"{name}: header {h!r} appears {len(idx)} times in {[c for c in header_row if c is not None]}"
            )
        out[h] = idx[0]
    return out


def parse_sheet(rows: list[tuple], spec: dict, name: str) -> pd.DataFrame:
    """One data sheet -> wide frame indexed by period_end, columns = the spec's header texts, values scaled.

    Rows after the header are either a quarter row (label in column A, consecutive with the previous one), or a row
    with nothing in the columns we read (blank lines, footnotes, the HE Revolving rows of the limit sheet).
    """
    unit = _cell(rows[1], 0) if len(rows) > 1 else None
    if _norm(unit) != _norm(spec["unit"]):
        raise ValueError(f"{name}: unit cell A2 should be {spec['unit']!r}, got {unit!r}")
    headers = [h for h, *_ in spec["columns"]]
    h = _header_row(rows, headers, name)
    pos = _column_positions(rows[h], headers, name)
    factor = float(spec["factor"])
    quarters: list[tuple[int, int]] = []
    values: list[dict[str, float]] = []
    for i, row in enumerate(rows[h + 1 :], start=h + 2):  # i is the 1-based row number in the sheet
        label = _cell(row, 0)
        try:
            quarter = parse_quarter(label)
        except ValueError:
            quarter = None
        if quarter is None:
            with_data = [hd for hd, j in pos.items() if _norm(_cell(row, j))]
            if with_data:
                raise ValueError(f"{name} row {i}: values under {with_data} but column A is {label!r}, not a quarter")
            continue
        if quarters and quarter != next_quarter(*quarters[-1]):
            prev = quarters[-1]
            raise ValueError(
                f"{name} row {i}: {quarter[0]}Q{quarter[1]} follows {prev[0]}Q{prev[1]}, quarters must be consecutive"
            )
        quarters.append(quarter)
        values.append({hd: _num(_cell(row, j), hd, name, i) * factor for hd, j in pos.items()})
    if not quarters:
        raise ValueError(f"{name}: no quarter rows after the header")
    index = pd.Index([pd.Timestamp(quarter_end(*q)) for q in quarters], name="period_end")
    return pd.DataFrame(values, index=index, columns=headers, dtype="float64")


def parse_workbook(path: Path, expected_quarter: tuple[int, int] | None = None) -> dict[str, pd.DataFrame]:
    """Workbook -> {sheet title: wide frame} for every sheet in SHEET_SPEC. Each must end on expected_quarter."""
    sheets = read_sheets(path)
    out = {}
    for spec in SHEET_SPEC:
        key = _norm(spec["title"])
        if key not in sheets:
            raise ValueError(f"{path.name}: no data sheet titled {spec['title']!r}. Titles found: {sorted(sheets)}")
        wide = parse_sheet(sheets[key], spec, f"{path.name} {spec['title']!r}")
        last = quarter_of(wide.index[-1].date())
        if expected_quarter is not None and last != expected_quarter:
            raise ValueError(
                f"{spec['title']!r}: last quarter in the sheet is {last[0]}Q{last[1]}, "
                f"the URL says {expected_quarter[0]}Q{expected_quarter[1]}"
            )
        out[spec["title"]] = wide
    return out


def to_facts(tables: dict[str, pd.DataFrame], pulled_at: str) -> pd.DataFrame:
    """Wide tables -> facts rows for every column in SHEET_SPEC, one per quarter (parse_sheet allows no gaps)."""
    frames = []
    for spec in SHEET_SPEC:
        wide = tables[spec["title"]]
        for header, metric, entity, entity_type in spec["columns"]:
            df = pd.DataFrame(
                {
                    "metric": metric,
                    "entity": entity,
                    "entity_type": entity_type,
                    "tier": "all",
                    "period_end": wide.index,
                    "period_type": "Q",
                    "value": wide[header].to_numpy(),
                    "source": SOURCE,
                    "pulled_at": pulled_at,
                }
            )
            frames.append(df[FACT_COLUMNS])
    return pd.concat(frames, ignore_index=True)


def parse_release(path: Path, pulled_at: str, expected_quarter: tuple[int, int] | None = None) -> pd.DataFrame:
    return to_facts(parse_workbook(path, expected_quarter), pulled_at)


def fetch(meta: pd.DataFrame, raw_dir: Path, session, pulled_at: str) -> pd.DataFrame:
    """Series come out of SHEET_SPEC, not out of meta. The loader rejects any that series.csv does not list."""
    today = dt.datetime.strptime(pulled_at[:10], "%Y-%m-%d").date()
    year, q, content = discover_latest(session, today)
    latest = raw_dir / "latest"
    latest.mkdir(parents=True, exist_ok=True)
    dest = latest / RAW_NAME
    dest.write_bytes(content)
    return parse_release(dest, pulled_at, expected_quarter=(year, q))
