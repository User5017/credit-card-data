"""NY Fed Household Debt and Credit fetcher: discovery walks back, the checked-in 2026 Q2 workbook parses, bad input fails.

The numbers asserted on the real workbook were read independently with openpyxl at hard-coded row and column
positions on 2026-09-07 (Page 3, 4, 10, 12, 13, 14 and 27 Data), not with the parser under test.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from carddash.fetchers import nyfed_hhdc as hhdc
from carddash.loader import validate
from carddash.schema import SERIES_KEY
from conftest import HHDC_FIXTURE, PULLED_AT

HTML_404 = (
    b'<!DOCTYPE html>\r\n<html lang="en">\r\n<head>\r\n<title>FEDERAL RESERVE BANK of NEW YORK</title></head>'
    b"<body></body></html>"
)
Q2_2026 = pd.Timestamp("2026-06-30")
N_LOAN_TYPE_QUARTERS = 94  # 2003Q1 .. 2026Q2
N_AGE_QUARTERS = 106  # 2000Q1 .. 2026Q2


def value(facts: pd.DataFrame, metric: str, period_end: str, entity: str = "CCP_ALL") -> float:
    sel = facts[
        (facts["metric"] == metric) & (facts["entity"] == entity) & (facts["period_end"] == pd.Timestamp(period_end))
    ]
    assert len(sel) == 1, f"{metric}/{entity}/{period_end}: {len(sel)} rows"
    return float(sel["value"].iloc[0])


# ---------- quarters and URLs ----------


@pytest.mark.parametrize(
    "label, quarter",
    [
        ("03:Q1", (2003, 1)),
        ("26:Q2", (2026, 2)),
        (" 11:Q4 ", (2011, 4)),
        (dt.datetime(2003, 3, 1), (2003, 1)),  # how the account sheet labels 2003Q1-2011Q1
        (dt.datetime(2010, 12, 1), (2010, 4)),
        (dt.date(2011, 3, 1), (2011, 1)),
    ],
)
def test_quarter_labels(label, quarter):
    assert hhdc.parse_quarter(label) == quarter


@pytest.mark.parametrize("bad", ["2026Q2", "2026:Q2", "26:Q5", "Return to Table of Contents", None, 3.5, ""])
def test_bad_quarter_labels_fail(bad):
    with pytest.raises(ValueError, match="not a quarter label"):
        hhdc.parse_quarter(bad)


def test_quarter_arithmetic_and_url():
    assert hhdc.quarter_end(2026, 2) == dt.date(2026, 6, 30)
    assert hhdc.quarter_end(2003, 1) == dt.date(2003, 3, 31)
    assert hhdc.quarter_of(dt.date(2026, 9, 7)) == (2026, 3)
    assert hhdc.previous_quarter(2026, 1) == (2025, 4)
    assert hhdc.next_quarter(2025, 4) == (2026, 1)
    assert hhdc.file_url(2026, 2) == (
        "https://www.newyorkfed.org/medialibrary/interactives/householdcredit/data/xls/HHD_C_Report_2026Q2.xlsx"
    )


# ---------- discovery ----------


class FakeSession:
    """Answers 200 for every URL, as the NY Fed's site does. The body says whether the file exists."""

    def __init__(self, available: dict[str, bytes]):
        self.available = available
        self.calls: list[str] = []

    def get(self, url):
        self.calls.append(url)
        body = self.available.get(url, HTML_404)

        class R:
            content = body
            status_code = 200

            def raise_for_status(self):
                pass

        return R()


def test_looks_like_xlsx_checks_the_zip_signature():
    assert hhdc.looks_like_xlsx(HHDC_FIXTURE.read_bytes()[:16])
    assert not hhdc.looks_like_xlsx(HTML_404)
    assert not hhdc.looks_like_xlsx(b"")


def test_discovery_walks_back_from_the_current_quarter():
    book = HHDC_FIXTURE.read_bytes()
    session = FakeSession({hhdc.file_url(2026, 2): book})
    year, q, content = hhdc.discover_latest(session, dt.date(2026, 9, 7))
    assert (year, q) == (2026, 2)
    assert content == book
    assert [u.rsplit("/", 1)[-1] for u in session.calls] == ["HHD_C_Report_2026Q3.xlsx", "HHD_C_Report_2026Q2.xlsx"]


def test_discovery_fails_loudly_when_nothing_is_found():
    with pytest.raises(ValueError, match="no HHD_C_Report workbook found for any of"):
        hhdc.discover_latest(FakeSession({}), dt.date(2026, 9, 7), max_back=3)


def test_fetch_end_to_end_with_a_fake_site(tmp_path, meta):
    session = FakeSession({hhdc.file_url(2026, 2): HHDC_FIXTURE.read_bytes()})
    raw = tmp_path / "nyfed_hhdc"
    facts = hhdc.fetch(meta, raw, session, "2026-09-07T12:00:00Z")
    assert (raw / "latest" / "HHD_C_Report.xlsx").read_bytes() == HHDC_FIXTURE.read_bytes()
    assert facts["period_end"].max() == Q2_2026
    assert len(facts) == 6 * N_LOAN_TYPE_QUARTERS + 6 * N_AGE_QUARTERS


# ---------- parsing the real release ----------


def test_sheet_shapes():
    tables = hhdc.parse_workbook(HHDC_FIXTURE, expected_quarter=(2026, 2))
    assert set(tables) == {s["title"] for s in hhdc.SHEET_SPEC}
    bal = tables["Total Debt Balance and Its Composition"]
    assert list(bal.columns) == ["Credit Card"] and len(bal) == N_LOAN_TYPE_QUARTERS
    assert bal.index[0] == pd.Timestamp("2003-03-31") and bal.index[-1] == Q2_2026
    age = tables["Transition into Serious Delinquency (90+) for Credit Cards by Age"]
    assert list(age.columns) == list(hhdc.AGE_GROUPS) and len(age) == N_AGE_QUARTERS
    assert age.index[0] == pd.Timestamp("2000-03-31") and age.index[-1] == Q2_2026
    accounts = tables["Number of Accounts by Loan Type"]  # date cells through 2011Q1, 'YY:Qn' labels after
    assert len(accounts) == N_LOAN_TYPE_QUARTERS
    assert accounts.loc["2003-03-31", "Credit Card"] == 469.81
    assert accounts.loc["2011-03-31", "Credit Card"] == 379.34
    assert accounts.loc["2011-06-30", "Credit Card"] == 389.17
    limit = tables["Credit Limit and Balance for Credit Cards and HE Revolving"]  # HE rows interleaved, skipped
    assert len(limit) == N_LOAN_TYPE_QUARTERS and limit.index.is_monotonic_increasing


def test_headline_values(hhdc_facts):
    f = hhdc_facts
    assert value(f, "hhdc_card_balances", "2026-06-30") == pytest.approx(1263.0)
    assert value(f, "hhdc_card_balances", "2026-03-31") == pytest.approx(1242.0)
    assert value(f, "hhdc_card_balances", "2025-06-30") == pytest.approx(1209.0)
    assert value(f, "hhdc_card_balances", "2011-03-31") == pytest.approx(696.4)
    assert value(f, "hhdc_card_balances", "2003-03-31") == pytest.approx(688.0)
    assert value(f, "hhdc_card_accounts", "2026-06-30") == 653.63
    assert value(f, "hhdc_card_limit", "2026-06-30") == pytest.approx(5559.0)
    assert value(f, "hhdc_card_limit", "2026-03-31") == pytest.approx(5474.0)
    assert value(f, "hhdc_card_limit", "2003-03-31") == pytest.approx(2550.0)
    assert value(f, "hhdc_card_dq90_rate_balances", "2026-06-30") == 12.92
    assert value(f, "hhdc_card_dq90_rate_balances", "2003-03-31") == 8.84
    assert value(f, "hhdc_card_transition_dq30", "2026-06-30") == 8.69
    assert value(f, "hhdc_card_transition_dq30", "2003-03-31") == 12.32
    assert value(f, "hhdc_card_transition_dq90", "2026-06-30") == 6.97
    assert value(f, "hhdc_card_transition_dq90", "2025-06-30") == 6.93
    assert value(f, "hhdc_card_transition_dq90", "2003-03-31") == pytest.approx(8.27348446019949)
    assert value(f, "hhdc_card_transition_dq90", "2026-06-30", "AGE:18-29") == pytest.approx(10.0786698080149)
    assert value(f, "hhdc_card_transition_dq90", "2026-06-30", "AGE:70+") == pytest.approx(6.343437351010106)
    assert value(f, "hhdc_card_transition_dq90", "2000-03-31", "AGE:18-29") == pytest.approx(9.631078849272566)
    # the 2026-08-11 press release arithmetic: +$21bn on the quarter, +$54bn on the year, +$85bn of limits
    bal = lambda pe: value(f, "hhdc_card_balances", pe)  # noqa: E731
    assert bal("2026-06-30") - bal("2026-03-31") == pytest.approx(21.0)
    assert bal("2026-06-30") - bal("2025-06-30") == pytest.approx(54.0)
    assert value(f, "hhdc_card_limit", "2026-06-30") - value(f, "hhdc_card_limit", "2026-03-31") == pytest.approx(85.0)


def test_facts_cover_every_series_in_series_csv_and_nothing_else(meta, hhdc_facts):
    mine = meta[meta["source"] == "nyfed_hhdc"]
    expected = set(mine[SERIES_KEY].itertuples(index=False, name=None))
    got = set(hhdc_facts[SERIES_KEY].drop_duplicates().itertuples(index=False, name=None))
    assert got == expected, f"missing {expected - got}, extra {got - expected}"
    assert len(expected) == 12  # six aggregate series, six age groups of the 90+ flow
    errors, warnings = validate(hhdc_facts, mine)
    assert errors == []
    assert warnings == []
    assert (hhdc_facts["period_type"] == "Q").all()
    assert (hhdc_facts["tier"] == "all").all()
    assert set(hhdc_facts.loc[hhdc_facts["entity"] == "CCP_ALL", "entity_type"]) == {"aggregate"}
    assert set(hhdc_facts.loc[hhdc_facts["entity"].str.startswith("AGE:"), "entity_type"]) == {"age"}
    counts = hhdc_facts.groupby(["metric", "entity"])["period_end"].count()
    is_age = counts.index.get_level_values("entity").str.startswith("AGE:")
    assert (counts[~is_age] == N_LOAN_TYPE_QUARTERS).all()
    assert (counts[is_age] == N_AGE_QUARTERS).all()
    assert len(hhdc_facts) == 6 * N_LOAN_TYPE_QUARTERS + 6 * N_AGE_QUARTERS


def test_limit_sheet_balance_column_tracks_the_balance_sheet():
    """The limit sheet repeats the card balance to fewer decimals; not loaded twice.

    The two columns differ only from 2009Q4 to 2012Q1 (nine quarters, up to $5bn) and agree exactly from 2012Q2 on.
    """
    sheets = hhdc.read_sheets(HHDC_FIXTURE)
    spec = {
        "title": "Credit Limit and Balance for Credit Cards and HE Revolving",
        "unit": "Trillions of $",
        "factor": hhdc.TRILLIONS_TO_BILLIONS,
        "columns": [("Credit Card Balance", "unused", "unused", "aggregate")],
    }
    p10 = hhdc.parse_sheet(sheets[hhdc._norm(spec["title"])], spec, "limit sheet")
    p3 = hhdc.parse_workbook(HHDC_FIXTURE)["Total Debt Balance and Its Composition"]
    assert p10.index.equals(p3.index)
    gap = (p10["Credit Card Balance"] - p3["Credit Card"]).abs()
    assert gap.max() <= 5.0 + 1e-6
    differs = gap[gap > 1e-6]
    assert len(differs) == 9
    assert differs.index.min() == pd.Timestamp("2009-12-31") and differs.index.max() == pd.Timestamp("2012-03-31")
    assert (gap[p3.index >= "2012-06-30"] <= 1e-6).all()


def test_file_must_end_on_the_quarter_the_url_names():
    with pytest.raises(ValueError, match="last quarter in the sheet is 2026Q2, the URL says 2026Q3"):
        hhdc.parse_workbook(HHDC_FIXTURE, expected_quarter=(2026, 3))


# ---------- format changes fail, on a small workbook in the report's layout ----------

QUARTERS = ["25:Q4", "26:Q1", "26:Q2"]


def write_workbook(path: Path, edit=None) -> Path:
    """Every sheet in SHEET_SPEC with three quarters, an unread 'Mortgage' column, a footnote, and the limit sheet's
    interleaved HE Revolving rows. `edit(title, rows)` may change a sheet before it is written."""
    wb = Workbook()
    wb.active.title = "TABLE OF CONTENTS"
    for n, spec in enumerate(hhdc.SHEET_SPEC, start=3):
        headers = [h for h, *_ in spec["columns"]]
        rows = [
            [spec["title"]],
            [spec["unit"], "Source: New York Fed Consumer Credit Panel/Equifax"],
            ["Return to Table of Contents"],
            [None, "Mortgage"] + headers,
        ]
        for k, q in enumerate(QUARTERS):
            rows.append([q, 100.0 + k] + [1.0 + k + 0.1 * c for c in range(len(headers))])
            if "Limit" in spec["title"]:
                rows.append([None, None] + [None] * len(headers) + [0.4 + k])  # HE Revolving, not read
        rows.append([])
        rows.append(["* footnote"])
        if edit:
            edit(spec["title"], rows)
        ws = wb.create_sheet(f"Page {n} Data")
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def test_synthetic_workbook_parses(tmp_path):
    facts = hhdc.parse_release(write_workbook(tmp_path / "ok.xlsx"), PULLED_AT, expected_quarter=(2026, 2))
    assert len(facts) == 12 * len(QUARTERS)
    assert value(facts, "hhdc_card_balances", "2026-06-30") == pytest.approx(3000.0)  # 3.0 trillion
    assert value(facts, "hhdc_card_limit", "2025-12-31") == pytest.approx(1000.0)
    assert value(facts, "hhdc_card_transition_dq90", "2025-12-31", "AGE:70+") == pytest.approx(1.5)


def _edit(fragment, fn):
    def edit(title, rows):
        if fragment in title:
            fn(rows)

    return edit


@pytest.mark.parametrize(
    "fragment, fn, match",
    [
        ("Total Debt Balance", lambda rows: rows[0].__setitem__(0, "Total Debt by Loan Type"), "no data sheet titled"),
        ("Total Debt Balance", lambda rows: rows[1].__setitem__(0, "Billions of $"), "unit cell A2 should be"),
        ("Number of Accounts", lambda rows: rows[3].__setitem__(2, "Cards"), "headers missing"),
        ("Number of Accounts", lambda rows: rows[3].__setitem__(1, "Credit Card"), "appears 2 times"),
        ("Percent of Balance 90", lambda rows: rows[5].__setitem__(2, "n/a"), "expected a number, got 'n/a'"),
        ("New Delinquent", lambda rows: rows[6].__setitem__(0, "26:Q3"), "2026Q3 follows 2026Q1, quarters must be consecutive"),
        ("by Age", lambda rows: rows[4].__setitem__(0, "26:Q1"), "quarters must be consecutive"),
        ("New Seriously", lambda rows: rows.insert(7, ["2026Q3", 1.0, 2.0]), "column A is '2026Q3', not a quarter"),
        ("New Seriously", lambda rows: rows.insert(7, ["Note", None, 2.0]), "values under \\['CC'\\]"),
    ],
)
def test_format_changes_fail_loudly(tmp_path, fragment, fn, match):
    path = write_workbook(tmp_path / "bad.xlsx", _edit(fragment, fn))
    with pytest.raises(ValueError, match=match):
        hhdc.parse_release(path, PULLED_AT, expected_quarter=(2026, 2))


def test_missing_data_sheets_fail(tmp_path):
    wb = Workbook()
    wb.active.title = "TABLE OF CONTENTS"
    wb.create_sheet("Chart3")
    wb.save(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError, match="no 'Page N Data' sheets"):
        hhdc.read_sheets(tmp_path / "empty.xlsx")
