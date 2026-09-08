"""NY Fed Credit Access Survey fetcher: the checked-in workbook parses, waves step by four months, the not-asked
marker is tolerated only as a leading gap, and format changes fail loudly.

The numbers asserted on the real workbook were read independently with openpyxl by sheet, row and column name on
2026-09-08, not with the parser under test.
"""

from __future__ import annotations

import datetime as dt

import pandas as pd
import pytest
import requests
from openpyxl import Workbook, load_workbook

from carddash.fetchers import nyfed_sce as sce
from carddash.loader import validate
from carddash.schema import SERIES_KEY
from conftest import PULLED_AT, REPO, SCE_FIXTURE

JUN_2026 = pd.Timestamp("2026-06-30")
N_WAVES = 39  # 2013 Q4 wave through June 2026, every four months
N_SCORE_WAVES = 38  # the credit score split starts one wave later, February 2014


def value(facts: pd.DataFrame, metric: str, entity: str, period_end: str) -> float:
    sel = facts[
        (facts["metric"] == metric) & (facts["entity"] == entity) & (facts["period_end"] == pd.Timestamp(period_end))
    ]
    assert len(sel) == 1, f"{metric}/{entity}/{period_end}: {len(sel)} rows"
    return float(sel["value"].iloc[0])


# ---------- waves ----------


@pytest.mark.parametrize("raw, wave", [("201310", (2013, 10)), (202606, (2026, 6)), (201402.0, (2014, 2)), (" 202510 ", (2025, 10))])
def test_parse_wave(raw, wave):
    assert sce.parse_wave(raw) == wave


@pytest.mark.parametrize("bad", ["2026-06", "202607", "20260630", "", None, "June 2026", 202613])
def test_bad_waves_fail(bad):
    with pytest.raises(ValueError):
        sce.parse_wave(bad)


def test_wave_arithmetic():
    assert sce.wave_end(2026, 6) == dt.date(2026, 6, 30)
    assert sce.wave_end(2013, 10) == dt.date(2013, 10, 31)
    assert sce.next_wave(2026, 6) == (2026, 10)
    assert sce.next_wave(2026, 10) == (2027, 2)
    assert sce.next_wave(2026, 2) == (2026, 6)


# ---------- the checked-in workbook ----------


def test_sheet_shapes():
    header, rows = sce.read_sheet(SCE_FIXTURE, "overall")
    assert header[0] == "date" and len(header) == 36
    assert len([r for r in rows if r[0]]) == N_WAVES
    header, rows = sce.read_sheet(SCE_FIXTURE, "demographics")
    assert len(header) == 26
    assert len([r for r in rows if r[0]]) == N_WAVES * 3 + N_SCORE_WAVES * 3
    with pytest.raises(ValueError, match="no sheet 'nope'"):
        sce.read_sheet(SCE_FIXTURE, "nope")


def test_headline_values(sce_facts):
    f = sce_facts
    # overall, June 2026 wave, read from the 'overall' sheet by column name
    assert value(f, "sce_card_application_rate", "SCE_ALL", "2026-06-30") == pytest.approx(29.89, abs=0.01)
    assert value(f, "sce_card_rejection_rate", "SCE_ALL", "2026-06-30") == pytest.approx(15.39, abs=0.01)
    assert value(f, "sce_card_limit_request_rate", "SCE_ALL", "2026-06-30") == pytest.approx(17.72, abs=0.01)
    assert value(f, "sce_card_limit_rejection_rate", "SCE_ALL", "2026-06-30") == pytest.approx(20.92, abs=0.01)
    assert value(f, "sce_any_rejection_rate", "SCE_ALL", "2026-06-30") == pytest.approx(16.13, abs=0.01)
    assert value(f, "sce_any_rejection_rate", "SCE_ALL", "2025-06-30") == pytest.approx(23.13, abs=0.01)
    assert value(f, "sce_observations", "SCE_ALL", "2026-06-30") == 895
    # by credit score, June 2026: the gap the whole panel exists to show
    assert value(f, "sce_any_rejection_rate", "SCORE:LT680", "2026-06-30") == pytest.approx(46.60, abs=0.01)
    assert value(f, "sce_any_rejection_rate", "SCORE:680-760", "2026-06-30") == pytest.approx(11.14, abs=0.01)
    assert value(f, "sce_any_rejection_rate", "SCORE:GE760", "2026-06-30") == pytest.approx(3.31, abs=0.01)
    assert value(f, "sce_discouraged_rate", "SCORE:LT680", "2026-06-30") == pytest.approx(20.74, abs=0.01)
    assert value(f, "sce_lender_closed_rate", "SCORE:LT680", "2026-06-30") == pytest.approx(20.92, abs=0.01)
    assert value(f, "sce_observations", "SCORE:LT680", "2026-06-30") == 151
    assert value(f, "sce_observations", "SCORE:GE760", "2026-06-30") == 508
    # by age
    assert value(f, "sce_any_rejection_rate", "AGE:GE60", "2026-06-30") == pytest.approx(6.07, abs=0.01)
    assert value(f, "sce_card_application_rate", "AGE:LE40", "2026-06-30") == pytest.approx(34.82, abs=0.01)


def test_shape_and_coverage(meta, sce_facts):
    f = sce_facts
    assert (f["period_type"] == "T").all() and (f["tier"] == "all").all()
    assert f["period_end"].max() == JUN_2026
    assert f["period_end"].min() == pd.Timestamp("2013-10-31")
    by_entity = f.groupby("entity")["period_end"].nunique()
    assert by_entity["SCE_ALL"] == N_WAVES
    assert by_entity["AGE:LE40"] == N_WAVES and by_entity["SCORE:LT680"] == N_SCORE_WAVES
    assert set(f.loc[f["entity"].str.startswith("SCORE:"), "entity_type"]) == {"score"}
    assert set(f.loc[f["entity"].str.startswith("AGE:"), "entity_type"]) == {"age"}
    assert set(f.loc[f["entity"] == "SCE_ALL", "entity_type"]) == {"aggregate"}
    # every wave is four months after the one before it, for every entity
    for _, grp in f.groupby(["metric", "entity"]):
        d = grp.sort_values("period_end")["period_end"]
        gaps = d.diff().dt.days.dropna()
        assert gaps.between(118, 124).all(), grp.name
    # the expectation question was not asked in the first wave and only there
    exp = f[f["metric"] == "sce_card_expected_rejection_rate"]
    assert len(exp) == N_WAVES - 1 and exp["period_end"].min() == pd.Timestamp("2014-02-28")
    # series.csv agrees
    mine = meta[meta["source"] == "nyfed_sce"]
    assert len(mine) == 52
    expected = set(mine[SERIES_KEY].itertuples(index=False, name=None))
    got = set(f[SERIES_KEY].drop_duplicates().itertuples(index=False, name=None))
    assert got == expected, f"missing {expected - got}, extra {got - expected}"
    errors, warnings = validate(f, mine)
    assert errors == []
    assert warnings == []


def test_percentages_are_percentages(sce_facts):
    rates = sce_facts[sce_facts["metric"] != "sce_observations"]
    assert rates["value"].between(0, 100).all()
    obs = sce_facts[sce_facts["metric"] == "sce_observations"]
    assert obs["value"].between(100, 2000).all()


# ---------- fetch end to end ----------


class FakeSession:
    def __init__(self, body: bytes | None = None, status: int = 200):
        self.body = body if body is not None else SCE_FIXTURE.read_bytes()
        self.status = status
        self.calls: list[str] = []

    def get(self, url):
        self.calls.append(url)
        body, status = self.body, self.status

        class R:
            content = body
            status_code = status

            def raise_for_status(self):
                if status >= 400:
                    raise requests.HTTPError(f"{status}")

        return R()


def test_fetch_end_to_end(tmp_path, meta, sce_facts):
    session = FakeSession()
    raw = tmp_path / "nyfed_sce"
    facts = sce.fetch(meta, raw, session, PULLED_AT)
    assert session.calls == [sce.FILE_URL]
    assert (raw / "latest" / sce.RAW_NAME).read_bytes() == SCE_FIXTURE.read_bytes()
    assert len(facts) == len(sce_facts)


def test_fetch_rejects_a_non_workbook(tmp_path, meta):
    with pytest.raises(ValueError, match="not an xlsx"):
        sce.fetch(meta, tmp_path / "nyfed_sce", FakeSession(b"<!DOCTYPE html><html>gone</html>"), PULLED_AT)
    with pytest.raises(requests.HTTPError):
        sce.fetch(meta, tmp_path / "nyfed_sce", FakeSession(status=503), PULLED_AT)


def test_looks_like_xlsx():
    assert sce.looks_like_xlsx(SCE_FIXTURE.read_bytes()[:8])
    assert not sce.looks_like_xlsx(b"<html>") and not sce.looks_like_xlsx(b"")


# ---------- format changes fail, on a small workbook in the survey's layout ----------

WAVES = [(2013, 10), (2014, 2), (2014, 6)]


def write_workbook(path, edit=None, waves=WAVES):
    """Both data sheets with three waves, an unread extra column and the attribution row."""
    wb = Workbook()
    wb.active.title = "Disclaimer"
    for sheet, columns in sce.SHEETS.items():
        ws = wb.create_sheet(sheet)
        headers = list(columns)
        ws.append(["Source: Survey of Consumer Expectations"])
        ws.append(["date", "group", "category", "Unused"] + headers)
        groups = [("all", "Overall")] if sheet == "overall" else [
            (g, c) for (g, c) in sce.ENTITIES if g != "all"
        ]
        rows = []
        for year, month in waves:
            for group, category in groups:
                rows.append([f"{year}{month:02d}", group, category, "x"] + [10.0 + i for i in range(len(headers))])
        if edit:
            edit(sheet, rows, headers)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def test_synthetic_workbook_parses(tmp_path):
    facts = sce.parse_release(write_workbook(tmp_path / "ok.xlsx"), PULLED_AT)
    assert facts["period_end"].tolist().count(pd.Timestamp("2014-06-30")) == len(sce.SHARED_COLUMNS) * 6 + len(sce.SHEETS["overall"])
    assert set(facts["entity"]) == {e for e, _ in sce.ENTITIES.values()}


def test_a_not_asked_marker_is_allowed_only_at_the_start(tmp_path):
    def leading(sheet, rows, headers):
        if sheet == "overall":
            rows[0][4 + headers.index("CCRejected")] = "N/A"

    facts = sce.parse_release(write_workbook(tmp_path / "leading.xlsx", leading), PULLED_AT)
    assert len(facts[facts["metric"] == "sce_card_rejection_rate"]) == 2  # the first wave is dropped

    def middle(sheet, rows, headers):
        if sheet == "overall":
            rows[1][4 + headers.index("CCRejected")] = "N/A"

    with pytest.raises(ValueError, match="has no value in 2014-02 after being asked from 2013-10"):
        sce.parse_release(write_workbook(tmp_path / "middle.xlsx", middle), PULLED_AT)

    def never(sheet, rows, headers):
        if sheet == "overall":
            for r in rows:
                r[4 + headers.index("CCRejected")] = "N/A"

    with pytest.raises(ValueError, match="is not asked in any wave"):
        sce.parse_release(write_workbook(tmp_path / "never.xlsx", never), PULLED_AT)


@pytest.mark.parametrize(
    "fn, match",
    [
        (lambda sheet, rows, headers: rows[0].__setitem__(1, "income"), "unknown group/category"),
        (lambda sheet, rows, headers: rows[0].__setitem__(0, "201312"), "the survey is fielded in"),
        (lambda sheet, rows, headers: rows[0].__setitem__(4, "n.d."), "expected a number"),
        (lambda sheet, rows, headers: rows[0].__setitem__(4, None), "expected a number"),
    ],
)
def test_bad_cells_fail(tmp_path, fn, match):
    with pytest.raises(ValueError, match=match):
        sce.parse_release(write_workbook(tmp_path / "bad.xlsx", fn), PULLED_AT)


def test_a_dropped_wave_fails(tmp_path):
    path = write_workbook(tmp_path / "gap.xlsx", waves=[(2013, 10), (2014, 6), (2014, 10)])
    with pytest.raises(ValueError, match="waves must step by four months"):
        sce.parse_release(path, PULLED_AT)


def test_a_missing_column_fails(tmp_path):
    path = write_workbook(tmp_path / "ok.xlsx")
    wb = load_workbook(path)
    ws = wb["demographics"]
    ws.cell(row=2, column=5, value="Renamed")
    wb.save(path)
    with pytest.raises(ValueError, match="columns missing or ambiguous"):
        sce.parse_release(path, PULLED_AT)


def test_history_must_start_at_the_first_wave(tmp_path):
    path = write_workbook(tmp_path / "late.xlsx", waves=[(2014, 2), (2014, 6), (2014, 10)])
    with pytest.raises(ValueError, match="history starts 2014-02-28, expected 2013-10-31"):
        sce.parse_release(path, PULLED_AT)


def test_golden_entries_trace_to_the_survey_page():
    text = (REPO / "checks" / "golden.yaml").read_text(encoding="utf-8")
    assert "sce_any_rejection_2025_06" in text and "sce_any_rejection_2026_06" in text
    assert "https://www.newyorkfed.org/microeconomics/sce/credit-access" in text
